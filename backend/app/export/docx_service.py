"""Export service — writes TipTap JSON content to .docx using python-docx.

Supports: paragraphs, headings (H1-H3), bold, italic, bullet lists, numbered lists.
"""

from __future__ import annotations

import io
import json
import logging
from typing import Any

from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

logger = logging.getLogger(__name__)

# TipTap node type -> docx heading level
_HEADING_LEVELS = {"heading": None, "h1": 1, "h2": 2, "h3": 3}


def _tiptap_to_docx(doc: DocxDocument, node: dict[str, Any]) -> None:
    """Recursively walk a TipTap JSON node and add content to the docx."""
    if not node:
        return

    node_type = node.get("type", "")

    # Document root — just walk children
    if node_type == "doc":
        for child in node.get("content", []) or []:
            _tiptap_to_docx(doc, child)
        return

    # Heading
    if node_type == "heading":
        level = node.get("attrs", {}).get("level", 1)
        level = max(1, min(3, level))
        text = _collect_text(node)
        doc.add_heading(text, level=level)
        return

    # Paragraph (the most common node)
    if node_type == "paragraph":
        text = _collect_text(node)
        p = doc.add_paragraph()
        _render_inline(doc, p, node)
        return

    # Bullet list
    if node_type == "bulletList":
        for item in node.get("content", []) or []:
            if item.get("type") == "listItem":
                _render_list_item(doc, item, numbered=False)
        return

    # Ordered list
    if node_type == "orderedList":
        for item in node.get("content", []) or []:
            if item.get("type") == "listItem":
                _render_list_item(doc, item, numbered=True)
        return

    # Blockquote
    if node_type == "blockquote":
        for child in node.get("content", []) or []:
            text = _collect_text(child)
            p = doc.add_paragraph(text)
            p.paragraph_format.left_indent = Pt(36)
            p.runs[0].italic = True if p.runs else False
        return

    # Code block
    if node_type == "codeBlock":
        text = _collect_text(node)
        p = doc.add_paragraph(text)
        if p.runs:
            p.runs[0].font.name = "Courier New"
            p.runs[0].font.size = Pt(10)
        return

    # Horizontal rule
    if node_type == "horizontalRule":
        doc.add_paragraph().add_run("———").bold = True
        return

    # Fallback: just extract text
    text = _collect_text(node)
    if text:
        doc.add_paragraph(text)


def _collect_text(node: dict[str, Any]) -> str:
    """Extract all text from a node (recursive)."""
    if not node:
        return ""
    parts: list[str] = []
    if node.get("type") == "text" and "text" in node:
        parts.append(node["text"])
    for child in node.get("content", []) or []:
        parts.append(_collect_text(child))
    return "".join(parts)


def _render_inline(doc: DocxDocument, p: Any, node: dict[str, Any]) -> None:
    """Render inline nodes (text with marks) into a paragraph."""
    for child in node.get("content", []) or []:
        if child.get("type") == "text":
            text = child.get("text", "")
            marks = child.get("marks", []) or []
            run = p.add_run(text)
            for mark in marks:
                mark_type = mark.get("type", "")
                if mark_type == "bold":
                    run.bold = True
                elif mark_type == "italic":
                    run.italic = True
                elif mark_type == "underline":
                    run.underline = True
                elif mark_type == "strike":
                    run.font.strike = True
                elif mark_type == "code":
                    run.font.name = "Courier New"
                    run.font.size = Pt(10)
        elif child.get("type") == "hardBreak":
            p.add_run().add_break()
        else:
            # nested inline — just grab text
            p.add_run(_collect_text(child))


def _render_list_item(doc: DocxDocument, item: dict[str, Any], numbered: bool) -> None:
    """Render a list item (can contain paragraphs)."""
    for child in item.get("content", []) or []:
        if child.get("type") == "paragraph":
            text = _collect_text(child)
            style = "List Number" if numbered else "List Bullet"
            p = doc.add_paragraph(text, style=style)
            _render_inline(doc, p, child)


def export_to_docx(content: dict[str, Any] | str) -> bytes:
    """Convert TipTap JSON content to .docx file bytes.

    Args:
        content: TipTap JSON as a dict or a JSON string.

    Returns:
        .docx file as bytes.
    """
    if isinstance(content, str):
        content = json.loads(content) if content else {}
    if not content:
        content = {"type": "doc", "content": [{"type": "paragraph"}]}

    doc = DocxDocument()
    _tiptap_to_docx(doc, content)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def import_from_docx(file_bytes: bytes) -> tuple[dict[str, Any], str, int]:
    """Read a .docx file and convert to TipTap JSON.

    Returns:
        Tuple of (content_json, title, word_count).
    """
    doc = DocxDocument(io.BytesIO(file_bytes))
    content_nodes: list[dict[str, Any]] = []
    title = ""
    word_count = 0

    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if not text:
            content_nodes.append({"type": "paragraph"})
            continue

        # First non-empty paragraph becomes the title
        if not title:
            title = text

        # Check heading style
        style_name = (para.style.name or "").lower() if para.style else ""

        if "heading 1" in style_name:
            content_nodes.append({
                "type": "heading", "attrs": {"level": 1},
                "content": [{"type": "text", "text": text}],
            })
        elif "heading 2" in style_name:
            content_nodes.append({
                "type": "heading", "attrs": {"level": 2},
                "content": [{"type": "text", "text": text}],
            })
        elif "heading 3" in style_name:
            content_nodes.append({
                "type": "heading", "attrs": {"level": 3},
                "content": [{"type": "text", "text": text}],
            })
        elif "list" in style_name:
            list_type = "orderedList" if "number" in style_name else "bulletList"
            # Try to append to last list if same type
            if content_nodes and content_nodes[-1].get("type") == list_type:
                content_nodes[-1]["content"].append({
                    "type": "listItem",
                    "content": [{"type": "paragraph", "content": [{"type": "text", "text": text}]}],
                })
            else:
                content_nodes.append({
                    "type": list_type,
                    "content": [{
                        "type": "listItem",
                        "content": [{"type": "paragraph", "content": [{"type": "text", "text": text}]}],
                    }],
                })
        else:
            # Regular paragraph — check for bold/italic runs
            inline_nodes = _docx_inline_to_tiptap(para)
            content_nodes.append({"type": "paragraph", "content": inline_nodes})

        word_count += len(text.split())

    if not content_nodes:
        content_nodes.append({"type": "paragraph"})

    return {"type": "doc", "content": content_nodes}, title, word_count


def _docx_inline_to_tiptap(para: Any) -> list[dict[str, Any]]:
    """Convert a python-docx paragraph's runs to TipTap inline nodes."""
    nodes: list[dict[str, Any]] = []
    for run in para.runs:
        text = run.text
        if not text:
            continue
        marks: list[dict[str, Any]] = []
        if run.bold:
            marks.append({"type": "bold"})
        if run.italic:
            marks.append({"type": "italic"})
        if run.underline:
            marks.append({"type": "underline"})
        node: dict[str, Any] = {"type": "text", "text": text}
        if marks:
            node["marks"] = marks
        nodes.append(node)
    if not nodes:
        nodes.append({"type": "text", "text": para.text})
    return nodes
